"""
Self-checking test for the per-field developer attributes (events / inline style / CSS class).

    python test_developer_attrs.py     # prints OK or raises AssertionError

They travel spreadsheet -> builder and back out through the spec export.
"""

import csv
import io
import os

from openpyxl import Workbook, load_workbook

from excel_importer import import_spreadsheet
from spreadsheet_exporter import _STD_HEADERS, export_spreadsheet


def _titled(de):
    return {"title": {"DE": de, "FR": de, "IT": de, "EN": de},
            "placeholder": {"DE": "", "FR": "", "IT": "", "EN": ""}}


def _field(comp_id, field_name, label, required=False, **dev):
    return {"id": comp_id, "fieldName": field_name, "required": required,
            "translations": _titled(label), **dev}


_PAYLOAD = {
    "formTitle": "Dev attrs",
    "deployment": {"baseName": "DevAttrs"},
    "sections": [
        {"id": "meta", "isMetadata": True,
         "components": [{"id": "meta-title", "translations": _titled("Dev attrs")}]},
        {"id": "S1", "translations": _titled("Angaben"), "components": [
            _field("text-input", "F_Lic", "Lizenz-Nummer", required=True,
                   onChange="vLicenceCheck(this);", inputStyle="max-width:20em",
                   inputClass="licenceNr"),
            _field("text-input", "F_Plain", "Ohne"),
            _field("number", "F_Num", "Anzahl", events='onchange="vNum(this)" onblur="trim(this)"'),
            _field("text-input", "F_Blur", "Nur Blur", events="onblur=\"check(this, 'x')\""),
            _field("phone", "F_Tel", "Telefon", required=True, inputStyle="max-width:12em;"),
            _field("phone", "F_Tel2", "Telefon 2"),
            _field("email", "F_Mail", "E-Mail"),
            _field("email", "F_Mailb", "E-Mail 2", required=True, inputClass="mail2"),
            _field("text-input", "F_Esc", "Escaping",
                   onChange="check(this, 'a', \"b\", ${x})",
                   inputStyle='font-family:"A B"', inputClass='good bad"x also-ok'),
        ]},
    ],
}


def _csv(header, rows):
    buf = io.StringIO()
    csv.writer(buf).writerows([header] + rows)
    return buf.getvalue().encode("utf-8")


_HEADER = ["Input ID", "Mussfeld PSPH", "Feld in PSPH deutsch", "Format PSPH",
           "Fachdatenbezeichner", "onChange (JS)", "Inline-Style (CSS)", "CSS-Klasse"]


def _run_import():
    raw = _csv(_HEADER, [
        ["1", "", "Angaben", "Header", "S_A", "", "", ""],
        ["2", "ja", "Lizenz-Nummer", "Text einzeilig", "F_Lic",
         "vLicenceCheck(this, 1);", "max-width:20em;", "licenceNr"],
        ["3", "", "Beleg", "Upload Feld einfach", "F_Beleg", "x(this)", "", ""],
        ["4", "", "Strasse", "Text einzeilig", "F_Str", "", "", ""],
        ["5", "", "Nr.", "Text einzeilig", "F_Nr", "", "width:4em", ""],
    ])
    payload = import_spreadsheet(raw, "form.csv")
    comps = {c.get("fieldName"): c for s in payload["sections"] for c in s["components"]}

    lic = comps["F_Lic"]
    assert lic["events"] == 'onchange="vLicenceCheck(this, 1);"', "the old column, as an onchange"
    assert lic["inputStyle"] == "max-width:20em;" and lic["inputClass"] == "licenceNr", lic
    assert "events" not in comps["F_Beleg"], "an upload has no row macro that takes it"

    ignored = {i["row"]: i["message"] for i in payload["issues"]
               if i["code"] == "developer-attrs-ignored"}
    assert set(ignored) == {4, 6}, ignored
    assert "only apply" in ignored[4], ignored[4]
    assert "merged" in ignored[6], ignored[6]

    print("OK: developer attributes import, and unusable ones are reported.")


def _run_export_round_trip():
    xlsx = export_spreadsheet(None, _PAYLOAD, author="t", note="n", version="1.0")
    back = {c.get("fieldName"): c
            for s in import_spreadsheet(xlsx, "spec.xlsx")["sections"] for c in s["components"]}
    assert back["F_Lic"]["events"] == 'onchange="vLicenceCheck(this);"', back["F_Lic"]
    assert back["F_Lic"]["inputStyle"] == "max-width:20em", back["F_Lic"]
    assert back["F_Num"]["events"] == 'onchange="vNum(this)" onblur="trim(this)"', back["F_Num"]
    assert "events" not in back["F_Plain"], back["F_Plain"]

    print("OK: developer attributes survive an export from scratch.")


def _legacy_workbook():
    """A PSPH sheet from before these columns existed."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Formular"
    ws.append(_STD_HEADERS[:-3])
    cols = {h: i for i, h in enumerate(_STD_HEADERS[:-3])}
    for label, fmt, fname in (("Angaben", "Header", "S_A"),
                              ("Lizenz-Nummer", "Text einzeilig", "F_Lic")):
        row = [""] * len(cols)
        row[cols["Feld in PSPH deutsch"]] = label
        row[cols["Format PSPH, Formate gem. TF 00"]] = fmt
        row[cols["Fachdatenbezeichner"]] = fname
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _run_overlay_adds_columns():
    source = _legacy_workbook()
    payload = import_spreadsheet(source, "legacy.xlsx")

    untouched = load_workbook(io.BytesIO(export_spreadsheet(source, payload, bump=False)))
    header = [c.value for c in untouched.active[1]]
    assert "Event (JS)" not in header, "no column is added until some field uses it"

    lic = next(c for s in payload["sections"] for c in s["components"]
               if c.get("fieldName") == "F_Lic")
    lic["events"] = 'onblur="vLicenceCheck(this);"'
    out = export_spreadsheet(source, payload, bump=False)

    header = [c.value for c in load_workbook(io.BytesIO(out)).active[1]]
    assert header[-1] == "Event (JS)", header
    assert "Inline-Style (CSS)" not in header, "only the column that is used"

    again = {c.get("fieldName"): c
             for s in import_spreadsheet(out, "legacy.xlsx")["sections"] for c in s["components"]}
    assert again["F_Lic"]["events"] == 'onblur="vLicenceCheck(this);"', again["F_Lic"]

    print("OK: an older workbook gains the column it needs on export.")


if __name__ == "__main__":
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    _run_import()
    _run_export_round_trip()
    _run_overlay_adds_columns()
