"""
Self-checking round-trip test for the PSPH spreadsheet exporter.

    python test_spreadsheet_exporter.py    # prints OK or raises AssertionError

Covers the Format PSPH mapping for rows *added* in the builder: a generated-from-scratch
workbook must re-import to the same component types it was exported from. Catalogs are
the interesting case — they come back as `select-multi` (identity is not carried in the
sheet) with their options and refIDs intact; a single checkbox keeps its identity via
its own `Kontrollkasten` format.

Also covers the deployment tab: the deployment ids the form sheet cannot express are
written there and read back on a fresh import, while a re-import keeps the ids the
builder already holds.
"""

import datetime
import io

from openpyxl import Workbook, load_workbook

from spreadsheet_exporter import export_spreadsheet, export_spreadsheet_versioned
from excel_importer import (import_spreadsheet, read_deployment_sheet,
                            DEPLOYMENT_FIELDS, DEPLOYMENT_SHEET)
from forms_store import merge_reimport
from json_builder import build_form_json, export_file_name, output_base_name


def _titled(de, fr=None, it=None, en=None):
    return {"title": {"DE": de, "FR": fr or de, "IT": it or de, "EN": en or de},
            "placeholder": {"DE": "", "FR": "", "IT": "", "EN": ""}}


def _opt(oid, de, ref=""):
    return {"optionId": oid, "fieldName": "F_" + oid, "refID": ref,
            "translations": {"DE": de, "FR": de, "IT": de, "EN": de}}


_PAYLOAD = {
    "formTitle": "Round Trip",
    "deployment": {"contactEmail": "kontakt@example.ch"},
    "sections": [
        {"id": "meta", "isMetadata": True,
         "components": [{"id": "meta-title", "translations": _titled("Round Trip")}]},
        {"id": "S1", "translations": _titled("Angaben"), "components": [
            {"id": "text-input", "fieldName": "F_Name", "translations": _titled("Name")},
            {"id": "select-single", "fieldName": "F_Single", "translations": _titled("Einfach"),
             "options": [_opt("s1", "A"), _opt("s2", "B")]},
            {"id": "select-multi", "fieldName": "F_Multi", "translations": _titled("Mehrfach"),
             "options": [_opt("m1", "Eins"), _opt("m2", "Zwei")]},
            {"id": "catalog-countries", "fieldName": "F_Staat", "required": True,
             "translations": _titled("Staat", "Etat", "Stato", "State"),
             "options": [_opt("c1", "Belgien", "BE"), _opt("c2", "Italien", "IT")]},
            {"id": "catalog-cantons", "fieldName": "F_Kanton",
             "translations": _titled("Kanton", "Canton", "Cantone", "Canton"),
             "options": [_opt("k1", "BL", "BL"), _opt("k2", "ZH", "ZH")]},
            {"id": "checkbox", "fieldName": "F_Grp", "translations": _titled("Frage?"),
             "options": [_opt("g1", "Eins"), _opt("g2", "Zwei")]},
            {"id": "single-checkbox", "fieldName": "F_Consent", "required": True,
             "translations": _titled("Ich stimme zu", "J'accepte", "Accetto", "I agree")},
        ]},
    ],
}


def _formats(xlsx):
    ws = load_workbook(io.BytesIO(xlsx))["Form"]
    header = [str(c.value) for c in ws[1]]
    fmt_col = next(i for i, h in enumerate(header) if "Format PSPH" in h)
    return [ws.cell(row=r, column=fmt_col + 1).value
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=fmt_col + 1).value]


def _run_catalog_round_trip():
    xlsx = export_spreadsheet(None, _PAYLOAD, author="test", note="initial", version="1.0")

    # Multi-select and both catalogs share the real sheets' multi spelling; single-select
    # keeps the "Drop Down Field" one.
    written = _formats(xlsx)
    assert written.count("Auswahlfeld mehrfach (Katalog)") == 3, written
    assert written.count("Auswahlfeld mehrfach Option") == 6, written
    assert written.count("Drop Down Field Titel") == 1, written
    assert written.count("Drop Down Field Option") == 2, written

    # A single checkbox is its own PSPH format, not a one-option Checkbox group: the group
    # renders the title in the left column *and* the option next to the box, which would
    # duplicate the label. Kontrollkasten keeps the generator's `desc=""`.
    assert written.count("Kontrollkasten") == 1, written
    assert written.count("Checkbox Titel") == 1, written
    assert written.count("Checkbox Option") == 2, written

    p = import_spreadsheet(xlsx, "round-trip.xlsx")
    assert not p["skipped"], p["skipped"]

    content = next(s for s in p["sections"] if not s.get("isMetadata"))
    ids = [c["id"] for c in content["components"]]
    assert ids == ["text-input", "select-single", "select-multi", "select-multi",
                   "select-multi", "checkbox", "single-checkbox"], ids

    multi = content["components"][2]
    assert [o["translations"]["DE"] for o in multi["options"]] == ["Eins", "Zwei"], multi["options"]

    # A catalog's refIDs are the catalog codes; losing them would break the generator.
    countries = content["components"][3]
    assert countries["required"] is True
    assert [(o["translations"]["DE"], o["refID"]) for o in countries["options"]] == [
        ("Belgien", "BE"), ("Italien", "IT")], countries["options"]

    cantons = content["components"][4]
    assert [o["refID"] for o in cantons["options"]] == ["BL", "ZH"], cantons["options"]

    consent = content["components"][6]
    assert consent["required"] is True
    assert not consent.get("options"), consent.get("options")
    assert consent["translations"]["title"]["FR"] == "J'accepte"

    print("OK: spreadsheet_exporter round-trip checks pass.")


def _run_dev_comment_overlay():
    """Fragen/Bemerkungen carries the builder's "Comments for devs" both ways.

    The overlay must leave the column alone for rows the builder has no comment field
    for — options above all, whose notes would otherwise be wiped on every export.
    """
    hdr = ["Input ID", "Mussfeld PSPH", "Feld in PSPH deutsch", "Name PSPH französisch",
           "Name PSPH italienisch", "Format PSPH, Formate gem. TF 00", "Fachdatenbezeichner",
           "RefID", "Complimentary ID", "Fragen/Bemerkungen"]
    rows = [
        ["", "", "Abschnitt", "Section", "Sezione", "Header", "sec1", "", "", ""],
        ["", "Ja", "Feld A", "Champ A", "Campo A", "Text einzeilig", "fA", "R1", "", "from the sheet"],
        ["", "", "Feld B", "Champ B", "Campo B", "Text einzeilig", "fB", "R2", "", ""],
        ["", "", "Wahl", "Choix", "Scelta", "Radiobutton Titel", "fW", "R3", "C3", "group note"],
        ["1", "", "Ja", "Oui", "Si", "Radiobutton Option", "oJa", "R4", "", "option note"],
    ]
    wb = Workbook()
    ws = wb.active
    ws.append(hdr)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    source = buf.getvalue()

    payload = import_spreadsheet(source, "src.xlsx")
    by_fn = {c.get("fieldName"): c
             for sec in payload["sections"] for c in sec["components"]}
    assert by_fn["fA"]["commentsForDevs"] == "from the sheet", by_fn["fA"]
    assert by_fn["fW"]["commentsForDevs"] == "group note", by_fn["fW"]

    by_fn["fA"]["commentsForDevs"] = "edited"
    by_fn["fB"]["commentsForDevs"] = "added"
    by_fn["fW"]["commentsForDevs"] = ""           # cleared in the builder

    out = export_spreadsheet(source, payload, author="test", note="n", bump=False)
    ws2 = load_workbook(io.BytesIO(out))[wb.active.title]
    got = {ws2.cell(row=r, column=7).value: ws2.cell(row=r, column=10).value
           for r in range(2, ws2.max_row + 1)}
    assert got["fA"] == "edited", got
    assert got["fB"] == "added", got
    assert got["fW"] is None, got                 # cleared through to the sheet
    assert got["oJa"] == "option note", got       # option row untouched

    print("OK: dev-comment overlay checks pass.")


def _run_deployment_tab():
    """The deployment ids ride along on their own tab, so the next fresh import has them."""
    deployment = {"compName": "permit", "usrId": "4711", "usrIdDev": "99",
                  "deId": "1001", "frId": "1002", "itId": "", "enId": "",
                  "contactEmail": "kontakt@example.ch", "validatorClass": "evd.V4723"}
    payload = dict(_PAYLOAD, deployment=deployment)
    xlsx = export_spreadsheet(None, payload, author="test", note="ids", version="1.0")

    wb = load_workbook(io.BytesIO(xlsx))
    assert DEPLOYMENT_SHEET in wb.sheetnames, wb.sheetnames
    # Only the filled ids are written; the blank ones leave their row empty to be filled in.
    stored = read_deployment_sheet(xlsx, "spec.xlsx")
    assert stored["usrId"] == "4711" and stored["usrIdDev"] == "99", stored
    assert stored["deId"] == "1001" and stored["frId"] == "1002", stored
    assert "itId" not in stored and "enId" not in stored, stored

    # A fresh import of that workbook needs no manual re-entry.
    imported = import_spreadsheet(xlsx, "spec.xlsx")
    assert imported["deployment"]["usrId"] == "4711", imported["deployment"]
    assert imported["deployment"]["deId"] == "1001", imported["deployment"]
    assert imported["deployment"]["validatorClass"] == "evd.V4723", imported["deployment"]
    assert imported["deployment"]["itId"] == "", imported["deployment"]

    # The form sheet stays the authority on the contact e-mail PSPH keeps there.
    assert imported["deployment"]["contactEmail"] == "kontakt@example.ch"

    # Exporting again updates the tab in place rather than appending a second set of rows.
    changed = dict(payload, deployment=dict(deployment, itId="1003", usrId="4712"))
    xlsx2 = export_spreadsheet(xlsx, changed, author="test", note="more ids", version="")
    ws = load_workbook(io.BytesIO(xlsx2))[DEPLOYMENT_SHEET]
    assert ws.max_row == 1 + len(DEPLOYMENT_FIELDS), ws.max_row  # header + one row per field
    again = read_deployment_sheet(xlsx2, "spec.xlsx")
    assert again["itId"] == "1003" and again["usrId"] == "4712", again

    # A workbook with no tab at all (an untouched legacy spec) imports as before.
    legacy = load_workbook(io.BytesIO(xlsx))
    del legacy[DEPLOYMENT_SHEET]
    buf = io.BytesIO()
    legacy.save(buf)
    assert read_deployment_sheet(buf.getvalue(), "spec.xlsx") == {}
    assert import_spreadsheet(buf.getvalue(), "spec.xlsx")["deployment"]["usrId"] == ""

    print("OK: deployment-tab round-trip checks pass.")


def _run_form_name_round_trip():
    """The name in My Forms comes back on a fresh import, not the Formulartitel heading."""
    payload = dict(_PAYLOAD, formTitle="Mein Gesuch 2026")
    back = import_spreadsheet(export_spreadsheet(None, payload), "spec.xlsx")
    assert back["formTitle"] == "Mein Gesuch 2026", back["formTitle"]
    assert "formName" not in back["deployment"], back["deployment"]

    overlay = export_spreadsheet(export_spreadsheet(None, _PAYLOAD), payload, bump=False)
    assert import_spreadsheet(overlay, "spec.xlsx")["formTitle"] == "Mein Gesuch 2026"
    print("OK: the form's name survives export and a fresh import.")


def _run_deployment_reimport():
    """Re-importing a corrected spec keeps the ids the builder already holds."""
    deployment = {"usrId": "4711", "deId": "1001"}
    xlsx = export_spreadsheet(None, dict(_PAYLOAD, deployment=deployment),
                              author="test", note="ids", version="1.0")
    fresh = import_spreadsheet(xlsx, "spec.xlsx")

    # The open form carries a corrected user id; the sheet still has the old one.
    stored = {"id": "f1", "title": "Round Trip",
              "payload": {"formTitle": "Round Trip", "sections": [],
                          "deployment": {"usrId": "9999", "deId": "", "frId": "1002"}}}
    merged = merge_reimport(stored, fresh)["payload"]["deployment"]

    assert merged["usrId"] == "9999", merged     # typed in the builder, so it wins
    assert merged["deId"] == "1001", merged      # blank in the builder, taken from the tab
    assert merged["frId"] == "1002", merged      # builder-only, survives the re-import

    print("OK: deployment re-import checks pass.")


def _run_output_base_name():
    """The generated artifacts are named from the Developer base name, not the form title."""
    assert output_base_name({}) == ''
    assert output_base_name({'baseName': '  '}) == ''
    assert output_base_name({'baseName': 'AbcQz356'}) == 'AbcQz356'
    # The pipeline prefixes "query" itself, so the file name as it should appear works too.
    assert output_base_name({'baseName': 'queryAbcQz356'}) == 'AbcQz356'
    assert output_base_name({'baseName': 'query_AbcQz356'}) == 'AbcQz356'
    assert output_base_name({'baseName': 'QueryAbcQz356'}) == 'AbcQz356'
    # A base name that is nothing but the prefix stays as typed rather than emptying out.
    assert output_base_name({'baseName': 'query'}) == 'query'
    assert output_base_name({'baseName': 'Abc qz/356'}) == 'Abc_qz_356'

    title = "EU Fahrtenheft fur den Gelegenheitsverkehr beantragen"
    plain = build_form_json(dict(_PAYLOAD, formTitle=title, deployment={}))
    assert plain['formName'] == 'EU_Fahrtenheft_fur_den_Gelegenheitsverkehr_beantragen'

    named = build_form_json(dict(_PAYLOAD, formTitle=title,
                                 deployment={'baseName': 'queryAbcQz356'}))
    # Both names drive the downstream naming (files and the identifiers pasted inside them).
    assert named['formName'] == 'AbcQz356', named['formName']
    assert named['formNameFull'] == 'AbcQz356', named['formNameFull']

    # It rides along on the deployment tab like the rest of the deployment settings, so the
    # next import of the workbook keeps generating the same artifact names.
    xlsx = export_spreadsheet(None, named_payload := dict(_PAYLOAD, formTitle=title,
                                                          deployment={'baseName': 'AbcQz356'}),
                              author="test", note="base name", version="1.0")
    assert read_deployment_sheet(xlsx, "spec.xlsx")["baseName"] == "AbcQz356"
    assert import_spreadsheet(xlsx, "spec.xlsx")["deployment"]["baseName"] == "AbcQz356"

    # The spec workbook keeps its own, title-derived name — that is the point of the setting.
    assert export_file_name(named_payload["formTitle"], "1.0", datetime.date(2026, 10, 23)) ==         "261023_EU_Fahrtenheft_fur_den_Gelegenheitsverkehr_beantragen_1.0"

    print("OK: output base-name checks pass.")


def _run_export_file_name():
    """The download name carries the date and the version the export actually recorded."""
    day = datetime.date(2026, 10, 22)
    assert export_file_name("EU Fahrtenheft für den Gelegenheitsverkehr beantragen", "2.1.1", day) ==         "261022_EU_Fahrtenheft_fuer_den_Gelegenheitsverkehr_beantragen_2.1.1"
    # Dots are the one piece of punctuation a version keeps; the rest folds to underscores.
    assert export_file_name("My form", " v2.1.1 ", day) == "261022_My_form_v2.1.1",         export_file_name("My form", " v2.1.1 ", day)
    assert export_file_name("My form", "2 (draft)", day) == "261022_My_form_2_draft"
    # Nothing recorded → the date and the name carry the file on their own.
    assert export_file_name("My form", "", day) == "261022_My_form"

    # A typed version is reported back as given...
    payload = {"formTitle": "My form", "sections": [
        {"id": "s1", "title": "S", "components": [dict(_titled("Feld"), **{
            "id": "c1", "type": "text", "fieldName": "fA"})]}]}
    first, typed = export_spreadsheet_versioned(None, payload, author="t", note="n", version="2.1")
    assert typed == "2.1", typed

    # ...and a blank one is bumped from the workbook's own history, which is why the name
    # has to be built from what comes back rather than from what was typed.
    _, bumped = export_spreadsheet_versioned(first, payload, author="t", note="n", version="")
    assert bumped == "2.2", bumped
    assert export_file_name("My form", bumped, day) == "261022_My_form_2.2"

    print("OK: export file-name checks pass.")


if __name__ == "__main__":
    _run_catalog_round_trip()
    _run_dev_comment_overlay()
    _run_deployment_tab()
    _run_deployment_reimport()
    _run_form_name_round_trip()
    _run_output_base_name()
    _run_export_file_name()
