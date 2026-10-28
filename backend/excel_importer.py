"""
Import a legacy PSPH form spreadsheet (.xlsx/.csv) into a formbuilder payload.

This is the inverse of json_builder's forward mapping: it reads the "Format PSPH"
column and rebuilds builder components, producing the same
{sections, formTitle, deployment, ...} shape the frontend emits — so an imported
form flows straight through the existing clean pipeline.

Phase 1 covers: the metadata header, common single-row inputs (text/number/date/
email/phone/textarea), option groups (radio/checkbox/dropdown), correspondence
language, uploads, info-text/header blocks and buttons. Adjacent Name+Vorname /
Strasse+Nr / PLZ+Ort text rows are recombined into split inputs. Catalogs and
show/hide-logic ID translation are reconstructed too. Aggregates (repeatable blocks)
are rebuilt as `isAggregate` sections: the rows between "Aggregat(e) start" and
"Aggregate end" become the block's components and the "Group N" rows its groups.

Columns are matched by header *name* (ASCII-folded, so umlaut/encoding variants
collapse together), never by position — different template versions reorder and
rename columns.

A row whose "Feld in PSPH deutsch" cell is empty is treated as Acta-Nova-managed
and skipped
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any, Optional

from field_names import ALLOWED_DESCRIPTION, is_valid_field_name
from input_events import format_events


# ── The 7 form-header metadata fields, in builder order (mirrors the frontend). ──
METADATA_FIELDS = [
    ('meta-title', 'Form Title'),
    ('meta-description', 'Description'),
    ('meta-department', 'Department'),
    ('meta-federalOffice', 'Federal Office'),
    ('meta-section', 'Section'),
    ('meta-address', 'Address'),
    ('meta-zipCity', 'ZIP and City'),
]

# id -> (palette name, material icon) for the components Phase 1 can emit.
_PALETTE = {
    'text-input': ('Text Input', 'text_fields'),
    'text-area': ('Text Area', 'notes'),
    'email': ('E-mail', 'mail'),
    'phone': ('Phone', 'phone'),
    'date': ('Date', 'calendar_today'),
    'number': ('Number', 'numbers'),
    'number-dropdown': ('Number Dropdown', 'pin'),
    'radio': ('Radio Button', 'radio_button_checked'),
    'checkbox': ('Checkbox Group', 'check_box'),
    'select-single': ('Dropdown', 'arrow_drop_down_circle'),
    'select-multi': ('Multi-Selection', 'playlist_add_check'),
    'korrespondenzsprache': ('Correspondence Language', 'translate'),
    'electronic-consent': ('Electronic Consent', 'gavel'),
    'upload-single': ('File Upload', 'upload_file'),
    'upload-multi': ('Multi-File Upload', 'cloud_upload'),
    'button': ('Button', 'smart_button'),
    'text': ('Text', 'short_text'),
}

# Format PSPH (ASCII-folded) -> builder component id, for single-row fields.
_SINGLE_FORMATS = {
    'texteinzeilig': 'text-input',
    'textmehrzeilig': 'text-area',
    'ganzzahl': 'number',
    'float': 'number',
    'festkommazahl': 'number',
    # Exact names only: a plain "Drop Down Field Titel" stays an option dropdown.
    'zahldropdown': 'number-dropdown',
    'numberdropdown': 'number-dropdown',
    'telefonnummer': 'phone',
    'datum': 'date',
    'uploadfeldeinfach': 'upload-single',
    'uploadfeldmehrfach': 'upload-multi',
    'button': 'button',
    'infobox': 'text',
    'infotext': 'text',
    'kontrollkasten': 'single-checkbox',
}

# The _SINGLE_FORMATS entries that import as a number accepting decimals.
_FLOAT_FORMATS = {'float', 'festkommazahl'}
_SPLIT_PAIRS = [
    (('name',),    ('vorname',),                              'name-vorname'),
    (('strasse',), ('nr', 'nummer', 'hausnr', 'hausnummer'), 'strasse-nr'),
    (('plz',),     ('ort',),                                  'plz-ort'),
]

# Split variant -> (palette name, material icon), mirroring the frontend palette.
_SPLIT_PALETTE = {
    'name-vorname': ('First Name + Name', 'badge'),
    'strasse-nr':   ('Street + No.', 'signpost'),
    'plz-ort':      ('ZIP + City', 'location_city'),
}

# Option-group title rows -> component id (each starts a group that consumes the
# option rows that follow it). Keyed on the canonical "title" spelling; _classify
# folds the German "Titel" to "title" first, so either spelling matches.
_TITLE_FORMATS = {
    'radiobuttontitle': 'radio',
    'checkboxtitle': 'checkbox',
    'dropdownfieldtitle': 'select-single',
}

# Option rows that attach to the current group.
_OPTION_FORMATS = {
    'radiobuttonoption', 'checkboxoption', 'dropdownfieldoption',
    'auswahlfeldeinfachoption', 'auswahlfeldmehrfachoption',
}

# ASCII-folded metadata formats -> meta id. The German spellings are the ones the PSPH
# template now writes; the English ones stay for sheets built against older templates.
# 'formulartitel' is keyed as 'formulartitle' because _classify folds Titel -> title first.
_METADATA_FORMATS = {
    'titleoftheform': 'meta-title',
    'descriptionoftheform': 'meta-description',
    'departmentoftheform': 'meta-department',
    'federalofficeoftheform': 'meta-federalOffice',
    'sectionoftheform': 'meta-section',
    'addressoftheform': 'meta-address',
    'zipandcityoftheform': 'meta-zipCity',
    'formulartitle': 'meta-title',
    'beschreibung': 'meta-description',
    'beschreibungformular': 'meta-description',
    'departement': 'meta-department',
    'bundesamt': 'meta-federalOffice',
    'sektionbundesamt': 'meta-section',
    'sektion': 'meta-section',
    'adressebundesamt': 'meta-address',
    'plzundortdesbundesamtes': 'meta-zipCity',
    'plzundortdesbundesamt': 'meta-zipCity',
}

# (column key, builder property, header the exporter writes). Passed straight through to the
# platform-side row macro, which only the single text-like inputs are rendered by.
DEVELOPER_ATTRS: tuple[tuple[str, str, str], ...] = (
    ('events', 'events', 'Event (JS)'),
    ('input_style', 'inputStyle', 'Inline-Style (CSS)'),
    ('input_class', 'inputClass', 'CSS-Klasse'),
)
DEVELOPER_ATTR_COMPONENTS = frozenset({'text-input', 'text-area', 'number', 'date', 'email', 'phone'})

# Canonical column key -> ASCII-folded header substrings that identify it.
_COLUMN_MATCHERS = {
    'input_id': ['inputid'],
    'show_logic': ['logikanzeigen'],
    'hide_logic': ['logikausblenden'],
    'hidden': ['verstecktesfeld'],
    'required': ['mussfeld'],
    'size': ['grsse', 'lnge', 'grsselnge'],
    'name_an_de': ['nameandeutsch'],
    'label_de': ['feldinpsphdeutsch'],
    'label_fr': ['namepsphfranz'],
    'label_it': ['namepsphitalien'],
    'label_en': ['namepsphengl'],
    'ph_it': ['hinweisimfeldpsphitalien'],
    'ph_fr': ['hinweisimfeldpsphfranz'],
    'ph_de': ['hinweisimfeldpsphdeutsch'],
    'mo_it': ['mouseoverpsphitalien'],
    'mo_fr': ['mouseoverpsphfranz'],
    'mo_de': ['mouseoverpsphdeutsch'],
    'format': ['formatpsph'],
    'field_name': ['fachdatenbezeichner'],
    'ref_id': ['refid'],
    'comp_id': ['complimentaryid'],
    'file_formats': ['dateiformate'],
    'comments': ['fragenbemerkungen'],
    'disable_mouse_over': ['mouseovernichtanzeigen'],
    'weblink': ['weblink'],
    'control': ['bemerkungsteuerung'],
    # 'onchange': sheets exported before events were generic call the column "onChange (JS)".
    'events': ['eventjs', 'onchange'],
    'input_style': ['inlinestyle'],
    'input_class': ['cssklasse', 'cssclass'],
}


def _fold(value: Any) -> str:
    """Lowercase, drop non-ASCII (so ö/ä and their mojibake collapse), keep a-z0-9."""
    text = '' if value is None else str(value)
    ascii_only = text.encode('ascii', 'ignore').decode('ascii')
    return re.sub(r'[^a-z0-9]', '', ascii_only.lower())


_LOGIC_TOKEN = re.compile(r'\d+(?:\.\d+)?|[\[\]]')

# These columns double as a notes column, and prose like "IF(30.06<Actual Date<=31.12)"
# was mined for Input IDs, inventing hide conditions on whatever rows shared those numbers.
_LOGIC_NOTATION = re.compile(r'^[0-9.,\[\]\s]+$')


def _flatten_or_group(node: list) -> list[str]:
    """Every Input ID token inside a (possibly nested) OR group, in order, de-duplicated."""
    out: list[str] = []
    for item in node:
        for tok in (_flatten_or_group(item) if isinstance(item, list) else [item]):
            if tok not in out:
                out.append(tok)
    return out


def _parse_logic_cell(raw: str) -> tuple[list, bool]:
    """Parse a "Logik anzeigen"/"Logik ausblenden" cell into AND/OR terms.

    The sheet notation (PSPH field guide, "Logik anzeigen") is a bracketed array whose
    top level is an AND — every entry must hold — while a nested array is an OR, of which
    at least one entry must hold: "[5, [8, 11]]" is 5 and either 8 or 11. Nesting below
    the first level stays OR, since the guide reads "[[5, 7, [11, 12]], 17]" as "5 or 7 or
    11 or 12, and 17", so a group collapses to a flat set of alternatives.

    Returns (terms, well_formed), where each term is an Input ID token (an AND entry) or a
    list of them (an OR group). A cell with no brackets at all is read as the top level,
    so a bare "5" and a bare "5, 8" mean 5, and 5 and 8 respectively. A cell that is not
    _LOGIC_NOTATION returns no terms, so the caller reports it unreadable.
    """
    if (raw or '').strip() and not _LOGIC_NOTATION.match(raw.strip()):
        return [], True
    stack: list[list] = [[]]
    well_formed = True
    for tok in _LOGIC_TOKEN.findall(raw or ''):
        if tok == '[':
            stack.append([])
        elif tok == ']':
            if len(stack) == 1:
                well_formed = False  # a ] that closes nothing
                continue
            stack[-2].append(stack.pop())
        else:
            stack[-1].append(tok)
    if len(stack) > 1:  # brackets left open
        well_formed = False
        while len(stack) > 1:
            stack[-2].append(stack.pop())

    top = stack[0]
    # The cell is normally wrapped in one pair of brackets, and that outer array *is* the
    # top-level AND list rather than a single OR group.
    if len(top) == 1 and isinstance(top[0], list):
        top = top[0]

    terms: list = []
    for item in top:
        if isinstance(item, list):
            alts = _flatten_or_group(item)
            if alts:
                terms.append(alts[0] if len(alts) == 1 else alts)
        else:
            terms.append(item)
    return terms, well_formed


def _as_bool(value: Any) -> bool:
    return _fold(value) in ('ja', 'yes', 'true', 'x', '1')


def _strip_tags(text: str) -> str:
    return re.sub(r'<[^>]+>', '', text or '')


def _strip_html(text: str) -> str:
    return _strip_tags(text).strip()


def _split_link(text: str) -> tuple[str, str]:
    """Return (label, href) for a value that may contain an <a href="...">…</a>.
    """
    if not text:
        return '', ''
    match = re.search(r'<a\b[^>]*href\s*=\s*"([^"]*)"[^>]*>(.*?)</a>', text, re.S | re.I)
    if not match:
        return _strip_html(text), ''
    href = (match.group(1) or '').strip()
    inner = _strip_html(match.group(2))
    before, after = _strip_tags(text[:match.start()]), _strip_tags(text[match.end():])
    if not inner:
        return f'{before}{after}'.strip(), href
    return f'{before}<a>{inner}</a>{after}'.strip(), href


def _decode_csv(raw: bytes) -> str:
    for encoding in ('utf-8-sig', 'utf-8', 'cp1252'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode('cp1252', 'replace')


def _read_rows(raw: bytes, filename: str) -> tuple[list[list[str]], str]:
    """Return the form sheet as string rows, plus the sheet's name.

    Rows start at spreadsheet row 1, so a row's index + 1 is the number the user sees
    in Excel — that is what makes an import issue findable in the source workbook.
    """
    if filename.lower().endswith('.csv'):
        text = _decode_csv(raw)
        return [[(c or '') for c in row] for row in csv.reader(io.StringIO(text))], ''

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    # A workbook usually has several sheets (the form, an internal Acta-Nova sheet,
    # a Version sheet, catalogs…). Several carry a "Format PSPH" column, so pick the
    # one with the most actual form rows rather than the first match.
    def sheet_rows(ws):
        ws.reset_dimensions()
        rows = [['' if c is None else str(c) for c in row]
                for row in ws.iter_rows(values_only=True)]
        width = max((len(row) for row in rows), default=0)
        for row in rows:
            row.extend([''] * (width - len(row)))
        return rows

    def form_score(rows):
        header = next((r for r in rows[:6] if any('formatpsph' in _fold(c) for c in r)), None)
        if header is None:
            return 0
        fmt_idx = _match_columns(header).get('format')
        if fmt_idx is None:
            return 0
        kinds = ('metadata', 'header', 'component', 'option', 'title')
        return sum(1 for r in rows
                   if fmt_idx < len(r) and _classify(r[fmt_idx])[0] in kinds)

    best = max(((sheet_rows(ws), ws.title) for ws in wb.worksheets),
               key=lambda pair: form_score(pair[0]), default=([], ''))
    wb.close()
    return best


def _match_columns(header: list[str]) -> dict[str, int]:
    folded = [_fold(h) for h in header]
    cols: dict[str, int] = {}
    for key, needles in _COLUMN_MATCHERS.items():
        for idx, h in enumerate(folded):
            if h and any(n in h for n in needles):
                cols[key] = idx
                break
    return cols


def _classify(fmt: str):
    """Return a classification tuple for a Format PSPH cell.

    Kinds: blank, internal, metadata, email-of-form, header, option, defer,
    aggregate-start, aggregate-group, aggregate-end,
    unknown, ('component', id, is_float) and ('title', id, silent_options) where
    `silent_options` means the group's option rows are rebuilt by the builder
    (electronic consent) and shouldn't be imported.
    """
    f = _fold(fmt)
    if not f:
        return ('blank',)
    # English "title" and German "Titel" are used interchangeably in these sheets
    # (and either spelling can appear in future templates), so collapse them.
    f = f.replace('titel', 'title')
    if f in _METADATA_FORMATS:
        return ('metadata', _METADATA_FORMATS[f])
    if 'emailoftheform' in f or 'emailbundesamt' in f:
        return ('email-of-form',)
    if f in ('header', 'textheader', 'sektion'):
        return ('header',)
    if 'keinfeld' in f and 'formularsprache' not in f:
        return ('internal',)  # "kein Feld (interne Daten)" — backend-only, not shown
    if f in _OPTION_FORMATS:
        return ('option',)
    if f in _TITLE_FORMATS:
        return ('title', _TITLE_FORMATS[f], False)
    if 'korrespondenzsprache' in f or 'formularsprache' in f:
        return ('title', 'korrespondenzsprache', False)
    if 'zustimmung' in f:
        return ('title', 'electronic-consent', True)  # fixed Yes/No, options rebuilt
    if 'auswahlfeld' in f or 'katalog' in f:
        single = 'einfach' in f and 'mehrfach' not in f
        return ('title', 'select-single' if single else 'select-multi', False)
    if f.startswith('group'):
        return ('aggregate-group',)
    if 'aggregat' in f:  # 'aggregat', so the German "Aggregat start/end" matches too
        if f.endswith('start'):
            return ('aggregate-start',)
        if f.endswith('end'):
            return ('aggregate-end',)
        return ('defer', 'unrecognised aggregate marker')
    if 'emailadresse' in f:
        return ('component', 'email', False)
    if f in _SINGLE_FORMATS:
        return ('component', _SINGLE_FORMATS[f], f in _FLOAT_FORMATS)
    return ('unknown', fmt)


class _Importer:
    def __init__(self, rows: list[list[str]], sheet_name: str = ''):
        self.rows = rows
        self.sheet_name = sheet_name
        self.skipped: list[str] = []
        # The same findings as `skipped`, but machine-readable: a stable code plus the
        # anchors (sheet row, label, Fachdatenbezeichner) the builder needs to decide
        # whether a finding has since been fixed by hand. See _issue.
        self.issues: list[dict] = []
        self.metadata: dict[str, dict[str, str]] = {}
        self.metadata_links: dict[str, str] = {}
        self.contact_email = ''
        self.sections: list[dict] = []
        self.current: Optional[dict] = None
        self.group: Optional[dict] = None  # last option-group component
        self.group_silent = False           # eat the group's option rows (fixed lists)
        self.aggregate: Optional[dict] = None  # open repeatable block, until "Aggregate end"
        self._n = 0
        self._row_no = 0  # 1-based spreadsheet row currently being read
        self._aggregate_row = 0
        self._optid_by_input_id: dict[str, str] = {}
        # (target dict, show terms, hide terms, source, mark-hidden) per row carrying logic,
        # resolved only once every option has been seen — see _resolve_pending_logic.
        self._pending_logic: list[tuple] = []
        # (normalised Input ID token, source label) for every show/hide reference seen,
        # checked against the complete option map at the end of the pass.
        self._logic_refs: list[tuple[str, str]] = []
        self._unreadable_logic: list[tuple[str, str, dict]] = []
        # Input ID -> (PSPH format, German label), to explain an unresolved reference.
        self._row_by_input_id: dict[str, tuple[str, str]] = {}
        self._dev_attr_row: dict[str, int] = {}

    def _uid(self, prefix: str) -> str:
        self._n += 1
        return f'{prefix}_{self._n}'

    def _issue(self, code: str, message: str, *, row: Optional[int] = None,
               label: str = '', field_name: str = '', input_id: str = '',
               target: Optional[dict] = None, skipped: bool = True) -> None:
        """Record one import finding, both as free text and as a structured issue.

        `code` is the stable identifier the builder maps to a resolution strategy;
        `label`/`field_name`/`target` are the anchors it matches against the canvas to
        decide whether the finding has since been fixed by hand. `skipped=False` keeps a
        finding about a row that *was* imported out of the skipped-rows list.
        """
        if skipped:
            self.skipped.append(message)
        issue: dict[str, Any] = {
            'code': code,
            'message': message,
            'row': self._row_no if row is None else row,
            'sheet': self.sheet_name,
        }
        if label:
            issue['label'] = label
        if field_name:
            issue['fieldName'] = field_name
        if input_id:
            issue['inputId'] = input_id
        if target:
            issue['target'] = target
        self.issues.append(issue)

    def _anchor(self, row, cols, kind: str) -> dict:
        """Where a row's content ended up in the builder, so an issue can be re-checked."""
        anchor: dict[str, Any] = {'kind': kind, 'row': self._row_no}
        label = self._get(row, cols, 'label_de')
        field_name = self._get(row, cols, 'field_name')
        if label:
            anchor['label'] = label
        if field_name:
            anchor['fieldName'] = field_name
        return anchor

    @staticmethod
    def _norm_input_id(token: str) -> str:
        """Normalise an Input ID / reference token so "[17]", "17" and "60.00" all match."""
        token = (token or '').strip().strip('[]').strip()
        try:
            return str(int(float(token)))
        except (TypeError, ValueError):
            return token

    def _defer_logic(self, target: dict, row, cols, source: str, *,
                     mark_hidden: bool = False, target_kind: str = 'field') -> None:
        """Parse this row's show/hide cells now; resolve their references later.

        Only the parse happens here, so a malformed cell is reported in document order.
        Resolution waits for the whole sheet (see _resolve_pending_logic). `mark_hidden`
        says a resolved Show/Hide condition implies the row starts hidden — true for
        fields, false for a section, whose heading carries that itself.
        """
        anchor = self._anchor(row, cols, target_kind)
        parsed_cells = []
        for key in ('show_logic', 'hide_logic'):
            raw = self._get(row, cols, key)
            terms, well_formed = _parse_logic_cell(raw)
            if raw.strip() and not terms:
                self._unreadable_logic.append((raw.strip(), source, anchor))
            elif terms and not well_formed:
                self._issue(
                    'logic-unbalanced',
                    f'{source} — show/hide condition "{raw.strip()}" has unbalanced '
                    f'brackets; the groups were read as best they could be, so check it',
                    label=anchor.get('label', ''), field_name=anchor.get('fieldName', ''),
                    target=anchor)
            parsed_cells.append(terms)
        if parsed_cells[0] or parsed_cells[1]:
            self._pending_logic.append(
                (target, parsed_cells[0], parsed_cells[1], source, mark_hidden, anchor))

    def _terms_to_conditions(self, terms: list, source: str, anchor: Optional[dict] = None):
        """Map parsed Input ID terms onto builder LogicCondition[] — an AND of entries.

        Every token is recorded (with `source` naming the field or section it belongs to)
        so genuinely dangling references can be reported once the sheet has been read.
        """
        out: list = []
        for term in terms:
            tokens = term if isinstance(term, list) else [term]
            oids: list[str] = []
            for tok in tokens:
                norm = self._norm_input_id(tok)
                if source:
                    self._logic_refs.append((norm, source, anchor))
                oid = self._optid_by_input_id.get(norm)
                if oid and oid not in oids:
                    oids.append(oid)
            if oids:
                out.append(oids if len(oids) > 1 else oids[0])
        return out or None

    def _resolve_pending_logic(self) -> None:
        """Resolve every show/hide cell against the complete option map.

        Resolving during the row walk silently dropped any reference to an option defined
        further *down* the sheet, because its Input ID wasn't in the map yet — and the
        dangling-reference report could not see it either, since by the end of the pass the
        id does exist. With the cell's top level being an AND, losing one entry weakens the
        condition and reveals the field earlier than the workbook specifies, so a forward
        reference like "[3, 8]" on row 5 would show the field on 3 alone.

        Runs before _merge_split_inputs, which copies the resolved values into the merged
        component.
        """
        for target, show_terms, hide_terms, source, mark_hidden, anchor in self._pending_logic:
            show = self._terms_to_conditions(show_terms, source, anchor)
            hide = self._terms_to_conditions(hide_terms, source, anchor)
            if show:
                target['showLogic'] = show
            if hide:
                target['hideLogic'] = hide
            if mark_hidden and (show or hide):
                target['hiddenField'] = True

    def run(self) -> dict:
        if not self.rows:
            return self._payload()
        header = self._find_header()
        cols = _match_columns(header)
        if 'format' not in cols:
            raise ValueError("Not a PSPH form sheet: no 'Format PSPH' column found.")
        self._has_input_id_column = 'input_id' in cols

        start = self.rows.index(header) + 1
        for offset, row in enumerate(self.rows[start:]):
            self._row_no = start + offset + 1  # 1-based, as Excel numbers it
            input_id = self._norm_input_id(self._get(row, cols, 'input_id'))
            if input_id:
                self._row_by_input_id[input_id] = (
                    self._get(row, cols, 'format'), self._get(row, cols, 'label_de'))
            self._row(row, cols)
            self._check_field_name(row, cols)
        if self.aggregate is not None:
            self._issue(
                'aggregate-unclosed',
                f'{self.aggregate["title"]} — aggregate has no "Aggregate end" row; every '
                f'field after its start was imported into the block',
                row=self._aggregate_row, label=self.aggregate['title'],
                field_name=self.aggregate.get('aggregateFachdatenbezeichner', ''))
        self._resolve_pending_logic()
        self._merge_split_inputs()
        self._report_unreadable_logic()
        self._report_unresolved_logic()
        self._report_unreachable_fields()
        self._report_unnamed_fields()
        self._drop_empty_unnamed()
        return self._payload()

    def _check_field_name(self, row, cols) -> None:
        field_name = self._get(row, cols, 'field_name')
        if is_valid_field_name(field_name):
            return
        label = self._get(row, cols, 'label_de')
        self._issue(
            'field-name-invalid',
            f'{label or field_name} — Fachdatenbezeichner "{field_name}" may only contain '
            f'{ALLOWED_DESCRIPTION}; rename it before generating',
            label=label, field_name=field_name)

    def _find_header(self) -> list[str]:
        for row in self.rows:
            if any('formatpsph' in _fold(c) for c in row):
                return row
        return self.rows[0]

    def _get(self, row, cols, key) -> str:
        idx = cols.get(key)
        return (row[idx].strip() if idx is not None and idx < len(row) and row[idx] else '')

    def _report_unresolved_logic(self) -> None:
        """Warn about show/hide references that match no imported option.

        Checked against the *complete* option map so backward references that were
        resolved during the pass stay silent; only genuinely dangling [Input ID]
        tokens (typo'd, or pointing at a row that wasn't imported) are reported.
        """
        if self._logic_refs and not self._row_by_input_id:
            self._report_logic_without_input_ids()
            return
        seen: set[tuple[str, str]] = set()
        for tok, source, anchor in self._logic_refs:
            if tok in self._optid_by_input_id or (tok, source) in seen:
                continue
            seen.add((tok, source))
            anchor = anchor or {}
            self._issue(
                'logic-dangling-ref',
                f'{source} — show/hide condition references [{tok}], '
                f'{self._describe_ref(tok)}. That term was dropped; the rest of the '
                f'condition was kept',
                row=anchor.get('row', 0), label=anchor.get('label', ''),
                field_name=anchor.get('fieldName', ''), input_id=tok,
                target=anchor or None)

    def _report_logic_without_input_ids(self) -> None:
        """Conditions point at options by Input ID; a sheet without any has nothing for them
        to point at. One finding naming that cause, rather than one per reference."""
        fields = list(dict.fromkeys(source for _, source, _ in self._logic_refs))
        examples = ', '.join(fields[:3]) + (f' and {len(fields) - 3} more' if len(fields) > 3 else '')
        cause = ('has no "Input ID" column' if not self._has_input_id_column
                 else 'leaves its "Input ID" column empty')
        self._issue(
            'logic-without-input-ids',
            f'This sheet {cause}, but {len(fields)} field(s) have show/hide conditions '
            f'({examples}). Conditions refer to options by Input ID, so they were dropped and '
            f'those fields always show. Add the Input IDs to the sheet and re-import, or set '
            f'the conditions in the builder',
            row=0, skipped=False)

    def _report_unreadable_logic(self) -> None:
        """Each distinct cell once with a row count: one repeated note is one finding."""
        grouped: dict[str, list[tuple[str, dict]]] = {}
        for raw, source, anchor in self._unreadable_logic:
            grouped.setdefault(raw, []).append((source, anchor or {}))
        for raw, hits in grouped.items():
            source, anchor = hits[0]
            others = len(hits) - 1
            also = f' (and {others} more row{"s" if others > 1 else ""})' if others else ''
            self._issue(
                'logic-unreadable',
                f'{source}{also} — show/hide condition "{raw}" is not an Input ID list '
                f'(expected e.g. "[5]" or "[5, [8, 11]]"); it was ignored rather than '
                f'guessed at, so the row keeps no condition from it',
                row=anchor.get('row', 0), label=anchor.get('label', ''),
                field_name=anchor.get('fieldName', ''), target=anchor or None)

    def _drop_empty_unnamed(self) -> None:
        self.sections = [s for s in self.sections
                         if s['components'] or s.get('title') or s.get('sourceRef')
                         or s.get('isAggregate')]

    def _report_unnamed_fields(self) -> None:
        """Generate needs a Fachdatenbezeichner for every field (the builder's
        missingFieldNames). One summary, not one per row: a sheet that leaves the column empty
        leaves it empty everywhere."""
        unnamed = []
        for section in self.sections:
            for comp in section['components']:
                if comp['id'] in ('text', 'electronic-consent'):
                    continue  # display text, and a consent the platform names itself
                sides = [(comp.get('splitConfig') or {}).get(s) or {} for s in ('a', 'b')]
                names = ([side.get('fieldName') for side in sides] if comp['id'] == 'split-input'
                         else [comp.get('fieldName')])
                if not all(names):
                    unnamed.append(self._title_de(comp) or comp.get('content') or comp['id'])
        if not unnamed:
            return
        examples = ', '.join(f'"{label}"' for label in unnamed[:3])
        more = f' and {len(unnamed) - 3} more' if len(unnamed) > 3 else ''
        self._issue(
            'field-names-missing',
            f'{len(unnamed)} field(s) have no Fachdatenbezeichner ({examples}{more}) — fill them in '
            f'before generating; each name is written into its row of the spreadsheet on export',
            row=0, skipped=False)

    def _report_unreachable_fields(self) -> None:
        """A field left "Verstecktes Feld = ja" with every term of its show condition
        unusable keeps the hidden flag with nothing to reveal it, so it can never appear.
        """
        for target, show_terms, _hide_terms, source, _mark, anchor in self._pending_logic:
            if not show_terms or target.get('showLogic') or not target.get('hiddenField'):
                continue
            anchor = anchor or {}
            self._issue(
                'logic-hidden-unreachable',
                f'{source} — hidden by default, and every term of its show condition was '
                f'unusable, so nothing is left to reveal it; as imported it cannot appear '
                f'in the form',
                row=anchor.get('row', 0), label=anchor.get('label', ''),
                field_name=anchor.get('fieldName', ''), target=anchor or None)

    def _describe_ref(self, tok: str) -> str:
        """Why a reference failed: usually the Input ID is real but is not an option row,
        and a condition can only key on an option.
        """
        row = self._row_by_input_id.get(tok)
        if row is None:
            return 'which no row in this sheet carries'
        fmt, label = row
        kind = _classify(fmt)[0] if fmt else ''
        named = f' ("{label[:40]}")' if label else ''
        if kind == 'title':
            return (f'the "{fmt}" group title{named} — reference the Input ID of the '
                    f'option to be chosen, not the question')
        if kind in ('component', 'metadata', 'header', 'unknown'):
            return (f'a "{fmt}" row{named}, not an option — only option rows '
                    f'(Radiobutton/Auswahlfeld/Checkbox Option) can be referenced')
        return f'a "{fmt}" row{named} that was not imported as an option'

    def _langs(self, row, cols, prefix) -> dict[str, str]:
        de = self._get(row, cols, f'{prefix}_de')
        fr = self._get(row, cols, f'{prefix}_fr')
        it = self._get(row, cols, f'{prefix}_it')
        en = self._get(row, cols, f'{prefix}_en')
        return {'DE': de, 'FR': fr, 'IT': it, 'EN': en}

    def _clean_langs_and_weblink(self, row, cols, prefix, keep_link=True) -> tuple[dict[str, str], str]:
        """Per-language labels plus the first href found across them.
        """
        cleaned = {}
        weblink = ''
        for lang, key in (('DE', 'de'), ('FR', 'fr'), ('IT', 'it'), ('EN', 'en')):
            label, href = _split_link(self._get(row, cols, f'{prefix}_{key}'))
            cleaned[lang] = label if keep_link else _strip_html(label)
            if href and not weblink:
                weblink = href
        return cleaned, weblink

    def _has_label(self, row, cols) -> bool:
        return any(self._langs(row, cols, 'label').values())

    def _has_de_label(self, row, cols) -> bool:
        """True when "Feld in PSPH deutsch" carries text — the PSPH-visible marker.

        Rows without it are Acta-Nova-managed (their German text lives in "Name AN
        deutsch") and are never applicant-facing, so they aren't builder components.
        """
        return bool(self._get(row, cols, 'label_de'))

    def _is_unnamed_section(self, row, cols) -> bool:
        """A header with no label in any language: a section left without a name. Acta-Nova
        blocks carry their text in "Name AN deutsch" instead, and stay skipped. A blank header
        that turns out to hold only Acta-Nova rows is dropped again (_drop_empty_unnamed)."""
        labels = ('label_de', 'label_fr', 'label_it', 'label_en', 'name_an_de')
        return not any(self._get(row, cols, key) for key in labels)

    def _row(self, row, cols):
        kind = _classify(self._get(row, cols, 'format'))
        tag = kind[0]

        if tag in ('blank', 'internal'):
            if tag == 'blank' and self._has_label(row, cols):
                label = self._get(row, cols, 'label_de') or '(row)'
                self._issue(
                    'format-missing',
                    f'{label} — empty "Format PSPH" cell; row skipped (set a format so it can be imported)',
                    label=self._get(row, cols, 'label_de'),
                    field_name=self._get(row, cols, 'field_name'))
            return
        if tag == 'metadata':
            keep_link = kind[1] == 'meta-description'
            titles, weblink = self._clean_langs_and_weblink(row, cols, 'label', keep_link)
            self.metadata[kind[1]] = titles
            if weblink:
                self.metadata_links[kind[1]] = weblink
            return
        if tag == 'email-of-form':
            candidates = [self._get(row, cols, k)
                          for k in ('label_de', 'label_fr', 'label_it', 'label_en', 'name_an_de')]
            email = next((c for c in candidates if '@' in c), '') or next((c for c in candidates if c), '')
            if email:
                self.contact_email = email
            return
        if tag == 'header':
            if not self._has_de_label(row, cols) and not self._is_unnamed_section(row, cols):
                # No "Feld in PSPH deutsch" → Acta-Nova-managed block (An-/Cc-Empfänger,
                # Versandbestätigung …). Not a PSPH section; skip it.
                self.group, self.group_silent = None, False
                return
            self._start_section(row, cols)
            return
        if tag == 'option':
            self._add_option(row, cols)
            return
        if tag == 'title':
            comp_id, silent = kind[1], kind[2]
            if not self._has_de_label(row, cols) and comp_id not in ('electronic-consent', 'korrespondenzsprache'):
                # No "Feld in PSPH deutsch" → Acta-Nova-managed group (Status,
                # Zuständige Sektion …). Skip it and eat its option rows.
                self.group, self.group_silent = None, True
                return
            self._add_component(comp_id, row, cols, is_float=False, is_group=True, silent_group=silent)
            return
        if tag == 'component':
            comp_id, is_float = kind[1], kind[2]
            if not self._has_de_label(row, cols):
                # No "Feld in PSPH deutsch" → Acta-Nova-managed field (Datum Gesuch,
                # Geschäfts-Nr. …). Not a PSPH form field; skip it (ends any open group).
                self.group, self.group_silent = None, False
                return
            self._add_component(comp_id, row, cols, is_float=is_float, is_group=False)
            return
        if tag == 'aggregate-start':
            self._start_aggregate(row, cols)
            return
        if tag == 'aggregate-group':
            self._add_aggregate_group(row, cols)
            return
        if tag == 'aggregate-end':
            self._end_aggregate(row, cols)
            return
        if tag == 'defer':
            label = self._get(row, cols, 'label_de') or self._get(row, cols, 'format')
            self._issue('format-deferred', f'{label or "(row)"} — {kind[1]}',
                        label=self._get(row, cols, 'label_de'),
                        field_name=self._get(row, cols, 'field_name'))
            self.group, self.group_silent = None, True  # eat any option rows that follow
            return
        # unknown
        if self._has_label(row, cols):
            self._issue(
                'format-unknown',
                f'{self._get(row, cols, "label_de") or "(row)"} — unknown format "{kind[1]}"',
                label=self._get(row, cols, 'label_de'),
                field_name=self._get(row, cols, 'field_name'))

    def _ensure_section(self):
        if self.current is None:
            self.current = {'id': self._uid('Section'), 'title': '', 'components': [], 'headerless': True}
            self.sections.append(self.current)

    def _start_section(self, row, cols):
        clean, _weblink = self._clean_langs_and_weblink(row, cols, 'label', keep_link=False)
        self.current = {
            'id': self._uid('Section'),
            'title': clean['DE'] or clean['FR'] or clean['IT'] or clean['EN'],
            'translations': {'title': clean},
            'components': [],
        }
        # Remember the header's Fachdatenbezeichner so a round-trip export can find and
        # update this exact header row in the original spreadsheet.
        source_ref = self._get(row, cols, 'field_name')
        if source_ref:
            self.current['sourceRef'] = source_ref
        else:
            # No Fachdatenbezeichner to find this header by again; the export uses its row.
            self.current['sheetRow'] = self._row_no
        ref_id = self._get(row, cols, 'ref_id')
        if ref_id:
            self.current['refID'] = ref_id
        comp_id = self._get(row, cols, 'comp_id')
        if comp_id:
            self.current['compId'] = comp_id
        if _as_bool(self._get(row, cols, 'hidden')):
            self.current['hiddenField'] = True
        ctx = f'section "{self.current["title"] or source_ref}"'
        self._defer_logic(self.current, row, cols, ctx, target_kind='section')
        self.sections.append(self.current)
        self.group, self.group_silent = None, False

    def _start_aggregate(self, row, cols):
        """Open an isAggregate section; rows up to "Aggregate end" become its components."""
        if self.aggregate is not None:  # PSPH sheets don't nest aggregates
            self._issue(
                'aggregate-autoclosed',
                f'{self.aggregate["title"]} — aggregate not closed before the next '
                f'"Aggregate start"; it was closed automatically',
                row=self._aggregate_row, label=self.aggregate['title'])
            self._end_aggregate(row, cols)
        clean, _weblink = self._clean_langs_and_weblink(row, cols, 'label', keep_link=False)
        section: dict[str, Any] = {
            'id': self._uid('Section'),
            'title': clean['DE'] or clean['FR'] or clean['IT'] or clean['EN'] or 'Aggregate',
            'translations': {'title': clean},
            'components': [],
            'isAggregate': True,
            'groups': [],
        }
        field_name = self._get(row, cols, 'field_name')
        if field_name:
            section['aggregateFachdatenbezeichner'] = field_name
            section['sourceRef'] = field_name
        ref_id = self._get(row, cols, 'ref_id')
        if ref_id:
            section['aggregateRefID'] = ref_id
        comp_id = self._get(row, cols, 'comp_id')
        if comp_id:
            section['aggregateCompId'] = comp_id
        self.sections.append(section)
        self.current = self.aggregate = section
        self._aggregate_row = self._row_no
        self.group, self.group_silent = None, False

    def _add_aggregate_group(self, row, cols):
        """Record a "Group N" row; the builder only counts them."""
        if self.aggregate is None:
            self._issue(
                'group-outside-aggregate',
                f'{self._get(row, cols, "format") or "(row)"} — "Group" row outside an aggregate',
                label=self._get(row, cols, 'label_de'),
                field_name=self._get(row, cols, 'field_name'))
            return
        self.aggregate['groups'].append({'groupId': self._uid('gr'), 'refID': ''})

    def _end_aggregate(self, row, cols):
        if self.aggregate is None:
            self._issue(
                'aggregate-end-orphan',
                f'{self._get(row, cols, "format") or "(row)"} — "Aggregate end" with no matching start')
            return
        self.aggregate = None
        self.current = None  # later rows start a fresh section
        self.group, self.group_silent = None, False

    def _add_option(self, row, cols):
        if self.group_silent:
            return  # electronic consent: its Yes/No options are rebuilt by the frontend
        if self.group is None:
            self._issue(
                'option-without-group',
                f'{self._get(row, cols, "label_de") or "(option)"} — option with no group',
                label=self._get(row, cols, 'label_de'),
                field_name=self._get(row, cols, 'field_name'),
                input_id=self._norm_input_id(self._get(row, cols, 'input_id')))
            return
        option_id = self._uid('opt')
        input_id = self._norm_input_id(self._get(row, cols, 'input_id'))
        if input_id:
            self._optid_by_input_id[input_id] = option_id
        self.group.setdefault('options', []).append({
            'optionId': option_id,
            'fieldName': self._get(row, cols, 'field_name'),
            'refID': self._get(row, cols, 'ref_id'),
            'translations': self._langs(row, cols, 'label'),
        })

    def _add_component(self, comp_id, row, cols, *, is_float, is_group, silent_group=False):
        self._ensure_section()
        name, icon = _PALETTE.get(comp_id, (comp_id, 'help'))
        labels, link_from_labels = self._clean_langs_and_weblink(row, cols, 'label')
        comp: dict[str, Any] = {
            'id': comp_id,
            'name': name,
            'icon': icon,
            'instanceId': self._uid(f'Field_{comp_id}'),
            'label': name,
            'required': _as_bool(self._get(row, cols, 'required')),
        }
        for key, prop in (('field_name', 'fieldName'), ('ref_id', 'refID'),
                          ('comp_id', 'compId'), ('comments', 'commentsForDevs'),
                          ('weblink', 'weblink')):
            val = self._get(row, cols, key)
            if val:
                comp[prop] = val
        if not comp.get('fieldName'):
            comp['sheetRow'] = self._row_no
        if link_from_labels and not comp.get('weblink'):
            comp['weblink'] = link_from_labels
        dev_attrs = {prop: val for key, prop, _ in DEVELOPER_ATTRS
                     if (val := self._get(row, cols, key))}
        if 'events' in dev_attrs:
            dev_attrs['events'] = format_events(dev_attrs['events'])
        if dev_attrs and comp_id in DEVELOPER_ATTR_COMPONENTS:
            comp.update(dev_attrs)
            self._dev_attr_row[comp['instanceId']] = self._row_no
        elif dev_attrs:
            self._issue(
                'developer-attrs-ignored',
                f'"{labels["DE"] or name}" — Event / Inline-Style / CSS-Klasse only apply to '
                f'text, number, date, e-mail and phone fields; ignored here',
                label=labels['DE'], field_name=self._get(row, cols, 'field_name'))
        if self._get(row, cols, 'hidden') and _as_bool(self._get(row, cols, 'hidden')):
            comp['hiddenField'] = True
        ctx = f'"{labels["DE"] or name}"'
        self._defer_logic(comp, row, cols, ctx, mark_hidden=True)
        if _as_bool(self._get(row, cols, 'disable_mouse_over')):
            comp['disableMouseOver'] = True
        if is_float:
            comp['allowDecimals'] = True
        size = self._get(row, cols, 'size')
        if size and size.isdigit():
            comp['maxLength'] = int(size)
        for key, prop in (('mo_de', 'mouseOverDe'), ('mo_fr', 'mouseOverFr'), ('mo_it', 'mouseOverIt')):
            val = self._get(row, cols, key)
            if val:
                comp[prop] = val

        if comp_id == 'text':
            comp['editField'] = 'content'
            comp['fontSize'] = 'base'
            comp['content'] = labels['DE']
            comp['translations'] = {'content': labels}
            if _fold(self._get(row, cols, 'format')) == 'infobox':
                comp['alignedInfobox'] = True
        else:
            comp['editField'] = 'title'
            comp['title'] = labels['DE']
            translations: dict[str, Any] = {'title': labels}
            placeholders = self._langs(row, cols, 'ph')
            if any(placeholders.values()):
                translations['placeholder'] = placeholders
                comp['placeholder'] = placeholders['DE']
            comp['translations'] = translations

        if comp_id == 'upload-single':
            fmts = self._get(row, cols, 'file_formats')
            if fmts:
                comp['fileFormats'] = fmts
        if comp_id == 'upload-multi':
            fmts = self._get(row, cols, 'file_formats')
            if fmts:
                comp['fileFormats'] = fmts

        if is_group:
            comp['options'] = []

        self.current['components'].append(comp)
        self.group = comp if is_group else None
        self.group_silent = silent_group

    @staticmethod
    def _title_de(comp: dict) -> str:
        return ((comp.get('translations') or {}).get('title') or {}).get('DE', '') or comp.get('title', '')

    @staticmethod
    def _split_variant_for(a_label: str, b_label: str) -> Optional[str]:
        fa, fb = _fold(a_label), _fold(b_label)
        for firsts, seconds, variant in _SPLIT_PAIRS:
            if fa in firsts and fb in seconds:
                return variant
        return None

    def _report_dropped_dev_attrs(self, c1: dict, c2: dict, variant: str) -> None:
        for comp in (c1, c2):
            if not any(comp.get(prop) for _, prop, _ in DEVELOPER_ATTRS):
                continue
            label = self._title_de(comp)
            self._issue(
                'developer-attrs-ignored',
                f'"{label}" — merged into one {variant} input, which takes no Event / '
                f'Inline-Style / CSS-Klasse; ignored here',
                row=self._dev_attr_row.get(comp['instanceId'], 0),
                label=label, field_name=comp.get('fieldName', ''))

    def _make_split(self, c1: dict, c2: dict, variant: str) -> dict:
        """Combine two adjacent single-row text fields into one split-input component."""
        name, icon = _SPLIT_PALETTE[variant]

        def side(c: dict) -> dict:
            tr = c.get('translations') or {}
            cfg: dict[str, Any] = {
                'titles': dict(tr.get('title') or {}),
                'placeholders': dict(tr.get('placeholder') or {}),
                'fieldName': c.get('fieldName', '') or '',
                'refID': c.get('refID', '') or '',
            }
            if c.get('sheetRow'):
                cfg['sheetRow'] = c['sheetRow']
            for prop in ('mouseOverDe', 'mouseOverFr', 'mouseOverIt'):
                if c.get(prop):
                    cfg[prop] = c[prop]
            return cfg

        split: dict[str, Any] = {
            'id': 'split-input',
            'name': name,
            'icon': icon,
            'instanceId': self._uid('Field_split'),
            'label': name,
            'required': bool(c1.get('required') or c2.get('required')),
            'splitVariant': variant,
            'splitConfig': {'a': side(c1), 'b': side(c2)},
            'title': self._title_de(c1),
        }
        # Shared single-value props live at the component level; prefer the first row.
        # refID stays per side (each row carries its own).
        for prop in ('compId', 'commentsForDevs', 'weblink'):
            val = c1.get(prop) or c2.get(prop)
            if val:
                split[prop] = val
        if c1.get('disableMouseOver') or c2.get('disableMouseOver'):
            split['disableMouseOver'] = True
        for prop in ('hiddenField', 'showLogic', 'hideLogic'):
            val = c1.get(prop) or c2.get(prop)
            if val:
                split[prop] = val
        return split

    def _merge_split_inputs(self) -> None:
        """Recombine adjacent Name+Vorname / Strasse+Nr / PLZ+Ort text rows into split inputs."""
        for section in self.sections:
            comps = section.get('components') or []
            merged: list[dict] = []
            i = 0
            while i < len(comps):
                c1 = comps[i]
                c2 = comps[i + 1] if i + 1 < len(comps) else None
                variant = None
                if c2 is not None and c1.get('id') == 'text-input' and c2.get('id') == 'text-input':
                    variant = self._split_variant_for(self._title_de(c1), self._title_de(c2))
                if variant:
                    self._report_dropped_dev_attrs(c1, c2, variant)
                    merged.append(self._make_split(c1, c2, variant))
                    i += 2
                else:
                    merged.append(c1)
                    i += 1
            section['components'] = merged

    def _payload(self) -> dict:
        sections: list[dict] = []
        if self.metadata or self.contact_email:
            components = []
            for i, (mid, label) in enumerate(METADATA_FIELDS):
                titles = self.metadata.get(mid, {'DE': '', 'FR': '', 'IT': '', 'EN': ''})
                components.append({
                    'id': mid, 'name': label, 'icon': 'edit_note',
                    'instanceId': f'Field_{mid}_{i}', 'label': label,
                    'required': False, 'editField': 'title',
                    'translations': {'title': titles},
                    'weblink': self.metadata_links.get(mid, '') or None,
                })
            sections.append({
                'id': self._uid('MetadataSection'), 'title': 'Form Header',
                'components': components, 'isMetadata': True,
            })
        sections.extend(self.sections)

        title_de = (self.metadata.get('meta-title') or {}).get('DE', '')
        deployment = {key: '' for key, _ in DEPLOYMENT_FIELDS if key != 'formName'}
        deployment['contactEmail'] = self.contact_email
        return {
            'formTitle': title_de or 'Imported form',
            'disableInputSize': False,
            'disableMouseOver': False,
            'deployment': deployment,
            'sections': sections,
            'skipped': self.skipped,
            'issues': self.issues,
        }


# ── Deployment tab ──────────────────────────────────────────────────
# The ids a package is deployed with (user ids, the per-language form ids, the validator
# class…) have no column in the PSPH form sheet — they are filled in in the builder. The
# exporter therefore creates them on their own tab, and a later fresh import of that workbook
# reads them back from here, so nobody has to type them a second time.

DEPLOYMENT_SHEET = 'Deployment'

DEPLOYMENT_FIELDS: tuple[tuple[str, str], ...] = (
    # The form's name in My Forms; not a deployment setting, but nowhere else in the sheet
    # keeps it (Formulartitel is the heading on the form, and often differs).
    ('formName', 'Formularname (Builder)'),
    ('baseName', 'Dateibasis (query...)'),
    ('compName', 'Komponentenname'),
    ('usrId', 'User ID (prod)'),
    ('usrIdDev', 'User ID (dev)'),
    ('contactEmail', 'Kontakt E-Mail'),
    ('deId', 'Formular-ID DE'),
    ('frId', 'Formular-ID FR'),
    ('itId', 'Formular-ID IT'),
    ('enId', 'Formular-ID EN'),
    ('wikiLink', 'Wiki-Seite'),
    ('issueLink', 'GitLab Issue-Nr.'),
    ('validatorClass', 'Validator-Klasse'),
)

DEPLOYMENT_HEADERS = ('Feld', 'Beschreibung', 'Wert')

_VALUE_HEADERS = ('wert', 'value')


def deployment_key(text: Any) -> Optional[str]:
    """The deployment field a tab row names, matched on the key or its label; None if neither."""
    folded = _fold(text)
    if not folded:
        return None
    for key, label in DEPLOYMENT_FIELDS:
        if folded in (_fold(key), _fold(label)):
            return key
    return None


def deployment_value_column(rows: list[list[str]]) -> int:
    """0-based value column of a deployment tab, from its header row (default: the third)."""
    for cells in rows:
        if deployment_key(cells[0] if cells else ''):
            continue
        for i, cell in enumerate(cells):
            if _fold(cell) in _VALUE_HEADERS:
                return i
    return 2


def read_deployment_sheet(raw: bytes, filename: str) -> dict[str, str]:
    """Deployment values stored on the workbook's deployment tab; {} when there is none.

    Only non-empty values are returned, so a half-filled tab leaves the rest alone.
    """
    if filename.lower().endswith('.csv'):
        return {}  # a single-sheet format has nowhere to carry the tab
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        return {}
    try:
        ws = next((s for s in wb.worksheets
                   if _fold(s.title) == _fold(DEPLOYMENT_SHEET)), None)
        if ws is None:
            return {}
        ws.reset_dimensions()  # same stale-<dimension> trap as _read_rows
        rows = [['' if c is None else str(c).strip() for c in row]
                for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    value_col = deployment_value_column(rows)
    found: dict[str, str] = {}
    for cells in rows:
        key = deployment_key(cells[0] if cells else '')
        if key and value_col < len(cells) and cells[value_col]:
            found[key] = cells[value_col]
    return found


def import_spreadsheet(raw: bytes, filename: str) -> dict:
    """Parse a PSPH .xlsx/.csv into a builder payload.

    Returns the {sections, formTitle, deployment, ...} payload plus `skipped`
    (human-readable lines) and `issues` — the same findings, coded and anchored to
    their sheet row so the builder can track which ones are still outstanding.

    Deployment ids come from the workbook's deployment tab when it has one, filling only
    what the form sheet itself cannot say — the contact e-mail stays the form sheet's,
    since that is where PSPH keeps it.
    """
    rows, sheet_name = _read_rows(raw, filename)
    payload = _Importer(rows, sheet_name).run()
    deployment = payload.get('deployment') or {}
    stored = read_deployment_sheet(raw, filename)
    if stored.get('formName'):
        payload['formTitle'] = stored.pop('formName')
    for key, value in stored.items():
        if not (deployment.get(key) or '').strip():
            deployment[key] = value
    payload['deployment'] = deployment
    return payload
