"""
Self-checking test: the exported workbook follows the builder's structure — deleted fields
leave the sheet, moved sections move — while rows the builder never shows stay put.

    python test_export_structure.py     # prints OK or raises AssertionError
"""

import csv
import io

from openpyxl import load_workbook

from excel_importer import import_spreadsheet
from spreadsheet_exporter import export_spreadsheet

_HEADER = ["Input ID", "Feld in PSPH deutsch", "Name AN deutsch", "Format PSPH",
           "Fachdatenbezeichner"]
_ROWS = [
    ["1", "Test", "", "Header", "S_Test"],
    ["2", "Text", "", "Text einzeilig", "F_Text1"],
    ["3", "Sprache", "", "Drop Down Field Titel", "F_Wahl"],
    ["4", "Eins", "", "Drop Down Field Option", ""],
    ["5", "Zwei", "", "Drop Down Field Option", ""],
    ["6", "", "Datum Gesuch", "Datum", "AN_Datum"],
    ["7", "", "", "Header", "S_Unnamed"],
    ["8", "Mail", "", "E-Mail-Adresse", "F_Email3"],
    ["9", "Text", "", "Text einzeilig", "F_Text2"],
    ["10", "", "An-Empfänger", "Header", "AN_Block"],
    ["11", "", "Empfänger", "Text einzeilig", "AN_Empf"],
]


def _source():
    buf = io.StringIO()
    csv.writer(buf).writerows([_HEADER] + _ROWS)
    return buf.getvalue().encode("utf-8")


def _names(xlsx):
    ws = load_workbook(io.BytesIO(xlsx)).active
    header = [c.value for c in ws[1]]
    name, label = header.index("Fachdatenbezeichner"), header.index("Feld in PSPH deutsch")
    return [r[name] or f"({r[label]})" for r in ws.iter_rows(min_row=2, values_only=True)
            if r[name] or r[label]]


def _payload():
    return import_spreadsheet(_source(), "form.csv")


def _sections(payload):
    return {s.get("sourceRef"): s for s in payload["sections"]}


def _run_unchanged():
    assert _names(export_spreadsheet(_source(), _payload(), bump=False)) == [
        "S_Test", "F_Text1", "F_Wahl", "(Eins)", "(Zwei)", "AN_Datum",
        "S_Unnamed", "F_Email3", "F_Text2", "AN_Block", "AN_Empf"]
    print("OK: an unedited form exports row for row.")


def _run_deleted_and_moved():
    payload = _payload()
    sections = _sections(payload)
    unnamed = sections["S_Unnamed"]
    unnamed["components"] = [c for c in unnamed["components"] if c["fieldName"] != "F_Email3"]
    test = sections["S_Test"]
    test["components"] = [c for c in test["components"] if c["fieldName"] != "F_Wahl"]
    payload["sections"] = [unnamed, test]

    assert _names(export_spreadsheet(_source(), payload, bump=False)) == [
        "S_Unnamed", "F_Text2",
        "S_Test", "F_Text1", "AN_Datum",
        "AN_Block", "AN_Empf",
    ], "deleted rows (and their options) go; the moved section moves; Acta-Nova rows stay"

    back = import_spreadsheet(export_spreadsheet(_source(), payload, bump=False), "out.xlsx")
    assert [(s["title"], [c["fieldName"] for c in s["components"]]) for s in back["sections"]] == [
        ("", ["F_Text2"]), ("Test", ["F_Text1"])]
    print("OK: deleted fields leave the sheet and moved sections move with them.")


def _run_duplicate_names_keep_order():
    rows = _ROWS + [["12", "Nochmal", "", "Text einzeilig", "F_Text1"]]
    buf = io.StringIO()
    csv.writer(buf).writerows([_HEADER] + rows)
    source = buf.getvalue().encode("utf-8")
    payload = import_spreadsheet(source, "form.csv")
    payload["sections"].reverse()
    names = _names(export_spreadsheet(source, payload, bump=False))
    assert names.index("S_Test") < names.index("S_Unnamed"), \
        "a name on two rows has no single place, so the sheet keeps its order"
    print("OK: sheets with a repeated Fachdatenbezeichner keep their row order.")


def _run_headers_without_name():
    """Headers imported without a Fachdatenbezeichner are updated in place, never duplicated."""
    rows = [["1", "Erster", "", "Header", "S_A"], ["2", "Eins", "", "Text einzeilig", "F_1"],
            ["3", "", "", "Header", ""], ["4", "Zwei", "", "Text einzeilig", "F_2"],
            ["5", "Ohne Name", "", "Header", ""], ["6", "Drei", "", "Text einzeilig", "F_3"]]
    buf = io.StringIO()
    csv.writer(buf).writerows([_HEADER] + rows)
    source = buf.getvalue().encode("utf-8")
    payload = import_spreadsheet(source, "form.csv")
    unnamed, labelled = payload["sections"][1], payload["sections"][2]
    assert unnamed["title"] == "" and unnamed["sheetRow"] == 4 and labelled["sheetRow"] == 6

    unnamed["title"], unnamed["translations"] = "Jetzt benannt", {"title": {"DE": "Jetzt benannt"}}
    labelled["sourceRef"] = "S_Typed"
    assert _names(export_spreadsheet(source, payload, bump=False)) == [
        "S_A", "F_1", "(Jetzt benannt)", "F_2", "S_Typed", "F_3"], \
        "the label lands in its own row, a typed name is written into it; no row is added"
    print("OK: headers without a Fachdatenbezeichner are found by their row.")


def _run_new_rows_in_sheet_without_names():
    """In a sheet with no Fachdatenbezeichner, new fields go to their section, not the top."""
    rows = [["", "Formular", "", "Title-of-the-form", ""],
            ["1", "Angaben", "", "Header", ""], ["2", "Name", "", "Text einzeilig", ""],
            ["3", "Gebühren", "", "Header", ""], ["4", "Mitteilungen", "", "Text mehrzeilig", ""]]
    buf = io.StringIO()
    csv.writer(buf).writerows([_HEADER] + rows)
    source = buf.getvalue().encode("utf-8")
    payload = import_spreadsheet(source, "form.csv")
    payload["sections"][1]["components"].append(
        {"id": "text-input", "fieldName": "F_Neu", "translations": {"title": {"DE": "Neu"}}})
    payload["sections"].append({"id": "S_new", "title": "", "sourceRef": "S_End", "components": [
        {"id": "date", "fieldName": "F_Datum", "translations": {"title": {"DE": "Datum"}}}]})

    assert _names(export_spreadsheet(source, payload, bump=False)) == [
        "(Formular)", "(Angaben)", "(Name)", "F_Neu", "(Gebühren)", "(Mitteilungen)",
        "S_End", "F_Datum"], "end of its own section; a new last section goes last"
    print("OK: new rows find their place in a sheet without Fachdatenbezeichner.")


def _run_fields_without_name():
    """Fields imported without a Fachdatenbezeichner are found by their row: edits land there,
    a name typed in the builder is written into it, and deleting the field removes it."""
    rows = [["1", "Angaben", "", "Header", ""], ["2", "Lizenz-Nummer", "", "Text einzeilig", ""],
            ["3", "Name", "", "Text einzeilig", ""], ["4", "Sprache", "", "Drop Down Field Titel", ""],
            ["5", "Deutsch", "", "Drop Down Field Option", ""],
            ["6", "Französisch", "", "Drop Down Field Option", ""],
            ["7", "Gebühren", "", "Header", ""], ["8", "Mitteilungen", "", "Text mehrzeilig", ""]]
    buf = io.StringIO()
    csv.writer(buf).writerows([_HEADER] + rows)
    source = buf.getvalue().encode("utf-8")
    payload = import_spreadsheet(source, "form.csv")
    licence, name, language = payload["sections"][0]["components"]
    assert (licence["sheetRow"], name["sheetRow"], language["sheetRow"]) == (3, 4, 5)

    licence["fieldName"] = "F_LizenzNr"
    name["translations"]["title"]["DE"] = "Name (neu)"
    payload["sections"][0]["components"] = [licence, name, {
        "id": "text-input", "fieldName": "F_Neu", "translations": {"title": {"DE": "Neu"}}}]

    assert _names(export_spreadsheet(source, payload, bump=False)) == [
        "(Angaben)", "F_LizenzNr", "(Name (neu))", "F_Neu", "(Gebühren)", "(Mitteilungen)"], \
        "named in place, relabelled in place, the deleted group and its options gone, no duplicates"
    print("OK: fields without a Fachdatenbezeichner are found by their row.")


if __name__ == "__main__":
    _run_fields_without_name()
    _run_new_rows_in_sheet_without_names()
    _run_headers_without_name()
    _run_unchanged()
    _run_deleted_and_moved()
    _run_duplicate_names_keep_order()
