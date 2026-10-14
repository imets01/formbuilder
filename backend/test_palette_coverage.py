"""
Self-checking guard that every palette component is mapped in both workbook layers.

    python test_palette_coverage.py     # prints OK or raises AssertionError

An unmapped component type fails silently twice: `_resolve_input_type` falls back to
"Text einzeilig", so it generates as a plain text field, and `_FORMAT_FOR` has no entry,
so `_write_row` leaves the Format cell empty and the row will not re-import. That is how
`single-checkbox`, `catalog-countries` and `catalog-cantons` sat broken in the palette.
The checklist in docs/extending.md exists to prevent it; this test enforces the part of
it that is mechanically checkable, by reading the ids straight out of the palette.
"""

import io
import os
import re

from openpyxl import load_workbook

from json_builder import COMPONENT_TYPE_MAP, OPTION_GROUP_FORMATS, build_form_json
from spreadsheet_exporter import _FORMAT_FOR, _OPTION_FORMAT_FOR, export_spreadsheet

_PALETTE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    'frontend', 'src', 'app', 'components', 'left-sidebar', 'left-sidebar.component.ts')

# Ids handled by a dedicated branch rather than a map entry.
_JSON_BRANCHES = {'split-input', 'text'}            # split sides / Infotext-vs-Header
_EXPORT_BRANCHES = {'split-input'}                  # emits one row per side
# The consent block is rendered by a backend macro off the `has_consent` flag, so it
# deliberately produces no row in either layer: build_form_json emits none and
# `_spec_for_component` returns None. Its OPTION_GROUP_FORMATS entry is never consulted.
_MACRO_DRIVEN = {'electronic-consent'}

_SILENT_FALLBACK = 'Text einzeilig'
_GENUINE_TEXT_INPUTS = {'text-input', 'split-input'}


def _palette_ids() -> list[str]:
    """The component ids offered in the left sidebar, in palette order, de-duplicated."""
    source = io.open(_PALETTE, encoding='utf-8').read()
    start = source.index('categories: ComponentCategory[] = [')
    end = source.index('\n  ];', start)
    ids: list[str] = []
    for match in re.finditer(r"\{\s*id:\s*'([^']+)'", source[start:end]):
        if match.group(1) not in ids:
            ids.append(match.group(1))
    assert len(ids) >= 15, f'palette parse looks wrong: {ids}'
    return ids


def _titled(de):
    return {'title': {'DE': de, 'FR': de, 'IT': de, 'EN': de},
            'placeholder': {'DE': '', 'FR': '', 'IT': '', 'EN': ''}}


def _one_field_payload(comp_id):
    comp = {'id': comp_id, 'fieldName': 'F_Probe', 'translations': _titled('Probe')}
    if comp_id in OPTION_GROUP_FORMATS:
        comp['options'] = [{'optionId': 'o1', 'fieldName': 'F_Opt', 'refID': '',
                            'translations': {'DE': 'Eins', 'FR': 'Un', 'IT': 'Uno', 'EN': 'One'}}]
    if comp_id == 'split-input':
        comp['splitVariant'] = 'name-vorname'
    return {'formTitle': 'Palette', 'deployment': {}, 'sections': [
        {'id': 'S', 'translations': _titled('Angaben'), 'components': [comp]}]}


def _run_maps_cover_palette():
    missing = []
    for comp_id in _palette_ids():
        in_json = (comp_id in COMPONENT_TYPE_MAP or comp_id in OPTION_GROUP_FORMATS
                   or comp_id in _JSON_BRANCHES)
        in_export = (comp_id in _FORMAT_FOR or comp_id in _EXPORT_BRANCHES
                     or comp_id in _MACRO_DRIVEN)
        has_option_format = (comp_id not in OPTION_GROUP_FORMATS
                             or comp_id in _OPTION_FORMAT_FOR
                             or comp_id in _MACRO_DRIVEN)
        if not in_json:
            missing.append(f'{comp_id}: no json_builder mapping (would generate as a text field)')
        if not in_export:
            missing.append(f'{comp_id}: no _FORMAT_FOR entry (would export without a Format)')
        if not has_option_format:
            missing.append(f'{comp_id}: no _OPTION_FORMAT_FOR entry (its options would be dropped)')
    assert not missing, 'palette components missing a workbook mapping:\n  ' + '\n  '.join(missing)
    print('OK: every palette component is mapped in both layers.')


def _run_each_type_generates_and_exports():
    """The maps can be present and still wrong, so drive the real code per component."""
    for comp_id in _palette_ids():
        payload = _one_field_payload(comp_id)
        form_json = build_form_json(payload)

        rows = [row for section in form_json['parts'].values()
                for key, row in section.items() if not key.startswith('__sectionhdr')]
        if comp_id in _MACRO_DRIVEN:
            assert not rows, f'{comp_id} is macro-driven but emitted rows: {rows}'
            assert form_json.get('has_consent'), f'{comp_id}: has_consent was not set'
        else:
            assert rows, f'{comp_id}: produced no form-structure row'
        if rows and comp_id not in _GENUINE_TEXT_INPUTS:
            assert rows[0]['input'] != _SILENT_FALLBACK, (
                f'{comp_id}: fell back to "{_SILENT_FALLBACK}" — it generates as a text field')

        xlsx = export_spreadsheet(None, payload, author='t', note='probe', version='1.0')
        ws = load_workbook(io.BytesIO(xlsx))['Form']
        header = [str(c.value) for c in ws[1]]
        fmt_col = next(i for i, h in enumerate(header) if 'Format PSPH' in h)
        label_col = next(i for i, h in enumerate(header) if h == 'Name AN deutsch')
        # Every row that carries a label must carry a Format, or it is the silent case:
        # written to the sheet, then skipped on the way back in. Checking all rows (not
        # just the component's own) also catches a group whose title lost its Format
        # while its option rows kept theirs.
        emitted = [(ws.cell(row=r, column=label_col + 1).value,
                    ws.cell(row=r, column=fmt_col + 1).value)
                   for r in range(2, ws.max_row + 1)
                   if ws.cell(row=r, column=label_col + 1).value
                   or ws.cell(row=r, column=fmt_col + 1).value]
        fields = [(label, fmt) for label, fmt in emitted
                  if fmt not in ('Header', 'Formulartitel')]
        if comp_id in _MACRO_DRIVEN:
            assert not fields, f'{comp_id} is macro-driven but wrote rows: {fields}'
            continue
        assert fields, f'{comp_id}: exported no row at all'
        if comp_id in OPTION_GROUP_FORMATS:
            assert len(fields) >= 2, (
                f'{comp_id}: exported {len(fields)} row(s) — an option group needs its title '
                f'row plus one row per option, so its choices survive the round trip')
        blank = [label for label, fmt in emitted if not fmt]
        assert not blank, (f'{comp_id}: row(s) {blank} exported with an empty Format — '
                           f'they will not re-import')

    print('OK: every palette component generates and exports with a Format.')


if __name__ == '__main__':
    _run_maps_cover_palette()
    _run_each_type_generates_and_exports()
