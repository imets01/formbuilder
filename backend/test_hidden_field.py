"""
Self-checking test: "hidden by default" survives a trip through the spec workbook.

    python test_hidden_field.py     # prints OK or raises AssertionError
"""

import io

from openpyxl import Workbook, load_workbook

from excel_importer import import_spreadsheet
from spreadsheet_exporter import export_spreadsheet
from spreadsheet_exporter import _STD_HEADERS
from test_developer_attrs import _titled

_PAYLOAD = {
    "formTitle": "Hidden",
    "sections": [
        {"id": "S1", "title": "Angaben", "sourceRef": "S_A", "translations": _titled("Angaben"),
         "components": [
             {"id": "text-input", "fieldName": "F_Hidden", "hiddenField": True,
              "translations": _titled("Versteckt")},
             {"id": "radio", "fieldName": "F_R", "hiddenField": True,
              "translations": _titled("Frage"),
              "options": [{"optionId": "o1", "translations": {"DE": "Ja"}},
                          {"optionId": "o2", "translations": {"DE": "Nein"}}]},
             {"id": "split-input", "splitVariant": "plz-ort", "hiddenField": True,
              "splitConfig": {"a": {"fieldName": "F_Plz", "titles": {"DE": "PLZ"}},
                              "b": {"fieldName": "F_Ort", "titles": {"DE": "Ort"}}}},
             {"id": "text-input", "fieldName": "F_Plain", "translations": _titled("Sichtbar")},
         ]},
        {"id": "S2", "title": "Teil 2", "sourceRef": "S_B", "hiddenField": True,
         "translations": _titled("Teil 2"),
         "components": [{"id": "text-input", "fieldName": "F_In2",
                         "translations": _titled("In zwei")}]},
    ],
}


def _imported(xlsx):
    payload = import_spreadsheet(xlsx, "spec.xlsx")
    sections = {s.get("sourceRef"): s for s in payload["sections"]}
    comps = {c.get("fieldName") or c.get("splitConfig", {}).get("a", {}).get("fieldName"): c
             for s in payload["sections"] for c in s["components"]}
    return payload, sections, comps


def _hidden_cells(xlsx):
    ws = load_workbook(io.BytesIO(xlsx)).active
    header = [c.value for c in ws[1]]
    hidden, name = header.index("Verstecktes Feld"), header.index("Fachdatenbezeichner")
    return {row[name]: row[hidden] for row in ws.iter_rows(min_row=2, values_only=True) if row[name]}


def _run_round_trip():
    xlsx = export_spreadsheet(None, _PAYLOAD)
    _, sections, comps = _imported(xlsx)
    for name in ("F_Hidden", "F_R", "F_Plz"):
        assert comps[name].get("hiddenField") is True, comps[name]
    assert not comps["F_Plain"].get("hiddenField")
    assert sections["S_B"].get("hiddenField") is True and not sections["S_A"].get("hiddenField")
    print("OK: hidden fields, groups, pairs and sections survive export and re-import.")


def _legacy_workbook():
    """A PSPH sheet without a Verstecktes Feld column, like the older PSPH templates."""
    headers = [h for h in _STD_HEADERS if h != "Verstecktes Feld"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Formular"
    ws.append(headers)
    for label, fmt, name in (("Angaben", "Header", "S_A"), ("Lizenz-Nummer", "Text einzeilig", "F_Lic")):
        row = [""] * len(headers)
        row[headers.index("Feld in PSPH deutsch")] = label
        row[headers.index("Format PSPH, Formate gem. TF 00")] = fmt
        row[headers.index("Fachdatenbezeichner")] = name
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _run_older_workbook():
    source = _legacy_workbook()
    payload, _, comps = _imported(source)
    assert "Verstecktes Feld" not in [c.value for c in load_workbook(io.BytesIO(source)).active[1]]

    untouched = export_spreadsheet(source, payload, bump=False)
    assert "Verstecktes Feld" not in [c.value for c in load_workbook(io.BytesIO(untouched)).active[1]], \
        "no column until some field is hidden"

    comps["F_Lic"]["hiddenField"] = True
    hidden = export_spreadsheet(source, payload, bump=False)
    assert _hidden_cells(hidden)["F_Lic"] == "Ja"
    _, _, back = _imported(hidden)
    assert back["F_Lic"].get("hiddenField") is True

    comps["F_Lic"]["hiddenField"] = False
    assert _hidden_cells(export_spreadsheet(hidden, payload, bump=False))["F_Lic"] is None, \
        "un-hiding clears the cell"
    print("OK: an older workbook gains the column, and un-hiding clears it.")


def _run_condition_not_stamped():
    payload = {"formTitle": "Logic", "sections": [{
        "id": "S1", "title": "Angaben", "sourceRef": "S_A", "translations": _titled("Angaben"),
        "components": [
            {"id": "radio", "fieldName": "F_Q", "translations": _titled("Frage"),
             "options": [{"optionId": "o_ja", "translations": {"DE": "Ja"}}]},
            {"id": "text-input", "fieldName": "F_Dep", "hiddenField": True,
             "showLogic": [["o_ja"]], "translations": _titled("Abhängig")},
        ]}]}
    cells = _hidden_cells(export_spreadsheet(None, payload))
    assert cells["F_Dep"] is None, "a show condition already implies hidden; the cell is not stamped"
    print("OK: fields with a condition are left to their condition.")


if __name__ == "__main__":
    _run_round_trip()
    _run_older_workbook()
    _run_condition_not_stamped()
